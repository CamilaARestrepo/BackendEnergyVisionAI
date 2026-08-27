# -*- coding: utf-8 -*-
import sys, os, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

WORD_FILE = r'resiudossolidos\Trabajo_v3.docx'
doc = Document(WORD_FILE)

p324 = doc.paragraphs[324]
p325 = doc.paragraphs[325]

print(f'[324] {p324.text[:80]}')
print(f'[325] {p325.text[:80]}')

# Replace placeholder
p325.clear()
desc = (
    'EnergyVision AI implementa una arquitectura basada en el patron Clean Architecture '
    'con cuatro capas: Presentacion, Aplicacion, Dominio e Infraestructura. '
    'El sistema sigue una arquitectura cliente-servidor donde el frontend (React 19 + Vite) '
    'se comunica con el backend (FastAPI + Uvicorn) mediante servicios REST sobre HTTP/JSON. '
    'La orquestacion del analisis se realiza mediante un grafo dirigido aciclico (DAG) '
    'implementado con LangGraph, con seis nodos secuenciales (validate, detection, waste, '
    'energy, enrichment, persist) y manejo de errores condicional con reintentos automaticos. '
    'La capa de datos utiliza SQLite con cifrado Fernet (AES-128) para proteger las claves '
    'API. El sistema soporta cuatro proveedores de IA seleccionables en tiempo de ejecucion: '
    'OpenAI (GPT-4o), Anthropic (Claude), Google Gemini y Ollama. '
    'A continuacion se presenta el diagrama general de la arquitectura del sistema:'
)
p325.add_run(desc)

# Insert caption + image after p325 (reverse order for correct sequence)
ref = p325._element

# Order wanted: p325, spacer, caption, image, spacer2
# Insert in reverse: spacer2, image, caption, spacer

sp2 = doc.add_paragraph('')
ref.addnext(sp2._element)

img_p = doc.add_paragraph()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = img_p.add_run()
img_path = os.path.join(os.path.dirname(WORD_FILE), 'diagrama_arquitectura_general.png')
run.add_picture(img_path, width=Inches(5.5))
ref.addnext(img_p._element)

cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap_run = cap.add_run('Ilustracion 7. Arquitectura General del sistema EnergyVision AI.')
cap_run.bold = True
cap_run.font.size = Pt(10)
cap_run.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
ref.addnext(cap._element)

sp1 = doc.add_paragraph('')
ref.addnext(sp1._element)

doc.save(WORD_FILE)
print('OK: guardado')

# Verify
d2 = Document(WORD_FILE)
ns = qn('w:drawing')
for i in range(323, 333):
    p = d2.paragraphs[i]
    has_img = False
    for r in p.runs:
        if r._element.findall(ns):
            has_img = True
            break
    t = p.text[:80] if p.text else '[empty]'
    tag = ' [IMG]' if has_img else ''
    print(f'  [{i}] {t}{tag}')
